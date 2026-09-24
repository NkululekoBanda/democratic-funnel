"""
File reading helpers and the raw-data inventory.

IEC / Stats SA files are messy:
  * CSVs may be UTF-16 (IEC 2011 results), UTF-8 with a BOM, or Latin-1, and may use ';' as delimiter
  * Excel reports often have title rows above the real header, and several sheets
These helpers handle all three.
"""
from pathlib import Path

import pandas as pd

from .config import RAW

ENCODINGS = ("utf-8-sig", "latin-1")   # utf-8-sig also reads plain UTF-8; latin-1 never fails
UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")
HEADER_KEYWORDS = ("municipal", "province", "ward", "registered", "votes", "party",
                   "voting district", "age", "code", "name", "population")


# ---------------------------------------------------------------- CSV
def _sniff_sep(path: Path, encoding: str) -> str:
    """Pick the delimiter that appears most often in the first line."""
    with open(path, encoding=encoding, errors="replace") as fh:
        first = fh.readline()
    counts = {sep: first.count(sep) for sep in (",", ";", "\t", "|")}
    return max(counts, key=counts.get)


def _encodings(path: Path) -> tuple:
    """Encodings to try, in order. UTF-16 is detected from its BOM first, because
    latin-1 'reads' a UTF-16 file without error but returns garbage column names."""
    with open(path, "rb") as fh:
        if fh.read(2) in UTF16_BOMS:
            return ("utf-16",)
    return ENCODINGS


def read_csv_any(path, **kwargs) -> pd.DataFrame:
    path = Path(path)
    last_err = None
    sep_arg = kwargs.pop("sep", None)
    for enc in _encodings(path):
        try:
            sep = sep_arg or _sniff_sep(path, enc)
            return pd.read_csv(path, encoding=enc, sep=sep, **kwargs)
        except UnicodeDecodeError as e:
            last_err = e
    raise ValueError(f"Could not decode {path}: {last_err}")


# ---------------------------------------------------------------- Excel
def find_header_row(path, sheet_name=0, scan_rows: int = 30,
                    keywords=HEADER_KEYWORDS) -> int:
    """Return the 0-based index of the first row that looks like a header:
    mostly filled, mostly text, and containing at least one keyword."""
    probe = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=scan_rows)
    if probe.empty:
        return 0
    width = probe.notna().sum(axis=1).max()
    for i, row in probe.iterrows():
        vals = row.dropna().astype(str).str.strip().str.lower()
        if len(vals) < max(2, 0.5 * width):
            continue
        texty = vals.str.contains(r"[a-z]", regex=True).mean() >= 0.6
        if texty and any(k in " ".join(vals) for k in keywords):
            return int(i)
    return 0


def read_excel_any(path, sheet_name=0, header="auto", **kwargs) -> pd.DataFrame:
    """Read an Excel sheet; header='auto' skips title rows above the real header."""
    if header == "auto":
        header = find_header_row(path, sheet_name=sheet_name)
    df = pd.read_excel(path, sheet_name=sheet_name, header=header, **kwargs)
    df = df.dropna(how="all").dropna(axis=1, how="all")          # blank rows/cols
    df.columns = [str(c).strip() for c in df.columns]
    return df


# ---------------------------------------------------------------- generic
def read_any(path, **kwargs) -> pd.DataFrame:
    """Read CSV or Excel with the robust defaults above."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        return read_excel_any(path, **kwargs)
    return read_csv_any(path, **kwargs)


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    """snake_case column names: 'Registered Voters ' -> 'registered_voters'."""
    out = df.copy()
    out.columns = (pd.Series(out.columns).astype(str)
                   .str.replace("\ufeff", "", regex=False)
                   .str.strip().str.lower()
                   .str.replace(r"[^a-z0-9]+", "_", regex=True)
                   .str.strip("_"))
    return out


# ---------------------------------------------------------------- inventory
def _excel_row_count(path: Path, sheet: str):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        n = wb[sheet].max_row
        wb.close()
        return n
    except Exception:
        return None


def inventory(folder=RAW, extensions=(".csv", ".xlsx", ".xlsm", ".xls")) -> pd.DataFrame:
    """One row per CSV file / per Excel sheet: size, rows, header row, columns."""
    folder = Path(folder)
    rows = []
    for f in sorted(folder.rglob("*")):
        if not f.is_file() or f.suffix.lower() not in extensions:
            continue
        base = {"file": str(f.relative_to(folder)).replace("\\", "/"),
                "size_kb": round(f.stat().st_size / 1024, 1)}
        try:
            if f.suffix.lower() == ".csv":
                head = read_csv_any(f, nrows=5)
                with open(f, "rb") as fh:
                    n = sum(1 for _ in fh) - 1
                rows.append({**base, "sheet": None, "header_row": 0, "rows": n,
                             "n_cols": head.shape[1],
                             "columns": " | ".join(map(str, head.columns))})
            else:
                for sheet in pd.ExcelFile(f).sheet_names:
                    h = find_header_row(f, sheet_name=sheet)
                    head = read_excel_any(f, sheet_name=sheet, header=h, nrows=5)
                    total = _excel_row_count(f, sheet)
                    rows.append({**base, "sheet": sheet, "header_row": h,
                                 "rows": (total - h - 1) if total else None,
                                 "n_cols": head.shape[1],
                                 "columns": " | ".join(map(str, head.columns))})
        except Exception as e:
            rows.append({**base, "sheet": None, "header_row": None, "rows": None,
                         "n_cols": None, "columns": f"ERROR: {e}"})
    return pd.DataFrame(rows, columns=["file", "sheet", "size_kb", "header_row",
                                       "rows", "n_cols", "columns"])
